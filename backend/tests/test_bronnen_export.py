from io import BytesIO

from openpyxl import load_workbook

from app.models import Batch, BronKandidaat, Company, Enrichment, ResearchRun


def test_export_bevat_gekozen_bron_en_geen_registervelden(client, db_session):
    batch = Batch(naam="bronwerkbank", jaar=2026, totaal=1)
    db_session.add(batch)
    db_session.flush()
    company = Company(batch_id=batch.id, naam="Voorbeeld", vestigingsnummer="V1")
    db_session.add(company)
    db_session.flush()
    run = ResearchRun(company_id=company.id, batch_id=batch.id, doel="bron", status="completed")
    db_session.add(run)
    db_session.flush()
    db_session.add(BronKandidaat(
        research_run_id=run.id, company_id=company.id,
        url="https://voorbeeld.nl/jaarverslag.pdf",
        canonical_url="https://voorbeeld.nl/jaarverslag.pdf",
        brontype="jaarverslag", status="geaccepteerd", wp_gevonden=42,
        eenheid="werkzame_personen", review_reason_code="juiste_bron_bruikbaar_bewijs",
    ))
    db_session.commit()

    response = client.get(f"/batches/{batch.id}/export.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    sheet = workbook.active
    rows = iter(sheet.iter_rows())
    headers = [cell.value for cell in next(rows)]
    values = [cell.value for cell in next(rows)]
    assert "Reviewredencode" in headers
    assert "WP goedgekeurd" not in headers
    assert values[headers.index("Bron-URL")] == "https://voorbeeld.nl/jaarverslag.pdf"
    assert values[headers.index("Gevonden waarde")] == 42


def test_export_noemt_altijd_de_website_ook_zonder_verrijking(client, db_session):
    """De website moet in elke rij staan, ook als er geen verrijking is.

    De kolom ontbrak helemaal in deze export. En beide exports schreven
    `enrichment.website_url if enrichment else website_url`, waardoor een
    verrijkingsrij zonder website — bijvoorbeeld na een mislukte Places-lookup —
    het aangeleverde adres wegdrukte.
    """
    batch = Batch(naam="website-export", jaar=2026, totaal=2)
    db_session.add(batch)
    db_session.flush()
    aangeleverd = Company(
        batch_id=batch.id, naam="Alleen aangeleverd", vestigingsnummer="V1",
        website_url="https://aangeleverd.nl",
    )
    gecorrigeerd = Company(
        batch_id=batch.id, naam="Ook verrijkt", vestigingsnummer="V2",
        website_url="https://oudnaam.nl",
    )
    db_session.add_all([aangeleverd, gecorrigeerd])
    db_session.flush()
    # Lege verrijking: mag het aangeleverde adres niet wegdrukken.
    db_session.add(Enrichment(company_id=aangeleverd.id, website_url=None))
    # Gevulde verrijking: gaat vóór, want die is tijdens de run gecorrigeerd.
    db_session.add(Enrichment(company_id=gecorrigeerd.id, website_url="https://nieuwnaam.nl"))
    db_session.commit()

    response = client.get(f"/batches/{batch.id}/export.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    headers = list(rows[0])
    per_naam = {rij[headers.index("Organisatienaam")]: rij for rij in rows[1:]}

    assert per_naam["Alleen aangeleverd"][headers.index("Website")] == "https://aangeleverd.nl"
    assert per_naam["Ook verrijkt"][headers.index("Website")] == "https://nieuwnaam.nl"
