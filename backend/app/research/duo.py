"""Rechtstreekse DUO-koppeling voor personeel in PO, VO en MBO."""
import asyncio
import csv
import html
import re
import unicodedata
from io import BytesIO, StringIO
from typing import Literal
from urllib.parse import urljoin

import requests
from openpyxl import load_workbook
from pydantic import BaseModel, Field, field_validator

from .query_planner import QueryContext
from .validation import SourceDocument


Sector = Literal["po", "vo", "mbo"]

_ADRES_URLS: dict[Sector, str] = {
    "po": "https://www.duo.nl/open_onderwijsdata/images/02.-alle-schoolvestigingen-basisonderwijs.csv",
    "vo": "https://duo.nl/open_onderwijsdata/images/02.-alle-vestigingen-vo.csv",
    "mbo": "https://duo.nl/open_onderwijsdata/images/01.-adressen-mbo-instellingen.csv",
}
_PERSONEEL_PAGINAS: dict[Sector, str] = {
    "po": "https://duo.nl/open_onderwijsdata/primair-onderwijs/personeel/onderwijspersoneel-po-in-aantal-personen.jsp",
    "vo": "https://duo.nl/open_onderwijsdata/voortgezet-onderwijs/personeel/in-aantal-personen.jsp",
    "mbo": "https://duo.nl/open_onderwijsdata/middelbaar-beroepsonderwijs/personeel/onderwijspersoneel-mbo-in-aantal-personen.jsp",
}
_PERSONEEL_MARKER: dict[Sector, str] = {
    "po": "onderwijspersoneel-po-in-personen",
    "vo": "onderwijspersoneel-vo-in-personen",
    "mbo": "onderwijspersoneel-mbo-in-personen",
}
_DOWNLOAD_CACHE: dict[str, bytes] = {}
_VESTIGING_CACHE: dict[Sector, list["DuoVestiging"]] = {}
_MAX_DOWNLOAD_BYTES = 80 * 1024 * 1024


class DuoVestiging(BaseModel):
    """Gevalideerde rij uit een officieel DUO-adresbestand."""

    sector: Sector
    instellingscode: str = Field(min_length=2, max_length=12)
    vestigingscode: str | None = None
    naam: str = Field(min_length=2)
    straat: str | None = None
    plaats: str | None = None
    gemeente: str | None = None
    provincie: str | None = None
    bevoegd_gezag: str | None = None

    @field_validator("instellingscode", "vestigingscode", mode="before")
    @classmethod
    def _normaliseer_code(cls, value):
        if value is None:
            return None
        return str(value).strip().upper()


class DuoPersoneelswaarde(BaseModel):
    """Gevalideerde personeelswaarde per DUO-instellingscode."""

    instellingscode: str = Field(min_length=2, max_length=12)
    jaar: int = Field(ge=2000, le=2100)
    aantal_personen: int | None = Field(default=None, ge=0)
    statistische_markering: str | None = None


def _norm(value: str | None) -> str:
    ascii_value = unicodedata.normalize("NFKD", value or "").encode(
        "ascii", "ignore",
    ).decode("ascii")
    return " ".join(re.findall(r"[a-z0-9]+", ascii_value.lower()))


def _naam_tokens(value: str | None) -> set[str]:
    stopwoorden = {
        "algemene", "basisschool", "college", "onderwijs", "school",
        "scholengemeenschap", "stichting", "van", "voor", "het", "de",
        "en", "rooms", "katholieke",
    }
    return {
        token for token in _norm(value).split()
        if len(token) >= 3 and token not in stopwoorden
    }


def _match_score(context: QueryContext, vestiging: DuoVestiging) -> float:
    gezocht = _naam_tokens(context.naam)
    gevonden = _naam_tokens(vestiging.naam)
    if not gezocht or not gevonden:
        return 0.0
    naam_score = len(gezocht & gevonden) / len(gezocht)
    if naam_score < (1.0 if len(gezocht) == 1 else 0.6):
        return 0.0

    score = naam_score * 0.75
    gemeente = set(_norm(context.gemeente).split())
    locatie = set(_norm(f"{vestiging.gemeente} {vestiging.plaats}").split())
    if gemeente and gemeente & locatie:
        score += 0.15
    if vestiging.straat and _norm(vestiging.straat) in _norm(context.adres):
        score += 0.10
    return score


def vind_vestigingen(
    context: QueryContext,
    vestigingen: list[DuoVestiging],
) -> list[DuoVestiging]:
    """Vind de beste naam-/gemeentematches en behoud groepsdeelinstellingen."""
    scores = [(item, _match_score(context, item)) for item in vestigingen]
    beste = max((score for _, score in scores), default=0.0)
    if beste < 0.6:
        return []
    return [item for item, score in scores if score >= beste - 0.08]


def lees_vestigingen(data: bytes, sector: Sector) -> list[DuoVestiging]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252")
    result = []
    for row in csv.DictReader(StringIO(text), delimiter=";"):
        naam = row.get("VESTIGINGSNAAM") or row.get("INSTELLINGSNAAM")
        if not naam or not row.get("INSTELLINGSCODE"):
            continue
        result.append(DuoVestiging(
            sector=sector,
            instellingscode=row["INSTELLINGSCODE"],
            vestigingscode=row.get("VESTIGINGSCODE") or None,
            naam=naam,
            straat=row.get("STRAATNAAM") or None,
            plaats=row.get("PLAATSNAAM") or None,
            gemeente=row.get("GEMEENTENAAM") or None,
            provincie=row.get("PROVINCIE") or None,
            bevoegd_gezag=row.get("BEVOEGD GEZAG NUMMER") or None,
        ))
    return result


def lees_personeelswaarden(
    data: bytes,
    instellingscodes: set[str],
    gevraagd_jaar: int | None,
) -> dict[str, DuoPersoneelswaarde]:
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    sheet = workbook["owtype-best-instelling"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    beschikbare_jaren = sorted(
        int(match.group(1))
        for header in headers
        if (match := re.fullmatch(r"PERSONEN (\d{4})", header))
    )
    kandidaten = [
        jaar for jaar in beschikbare_jaren
        if gevraagd_jaar is None or jaar <= gevraagd_jaar
    ]
    if not kandidaten:
        workbook.close()
        return {}
    jaar = max(kandidaten)
    code_index = headers.index("INSTELLINGSCODE")
    waarde_index = headers.index(f"PERSONEN {jaar}")
    result = {}
    for row in rows:
        code = str(row[code_index] or "").strip().upper()
        if code not in instellingscodes:
            continue
        raw = row[waarde_index]
        aantal = None
        marker = None
        if isinstance(raw, (int, float)):
            aantal = max(0, round(raw))
        elif raw not in {None, ""}:
            marker = str(raw).strip()
        result[code] = DuoPersoneelswaarde(
            instellingscode=code,
            jaar=jaar,
            aantal_personen=aantal,
            statistische_markering=marker,
        )
    workbook.close()
    return result


def _scope(
    sector: Sector,
    matches: list[DuoVestiging],
    alle_vestigingen: list[DuoVestiging],
) -> str:
    codes = {item.instellingscode for item in matches}
    locaties = [
        item for item in alle_vestigingen if item.instellingscode in codes
    ]
    if sector == "mbo":
        return "instelling"
    if len(codes) == 1 and len(locaties) == 1:
        return "vestiging"
    if locaties and all(_norm(item.provincie) == "limburg" for item in locaties):
        return "limburg"
    return "instelling"


def bouw_duo_bron(
    context: QueryContext,
    matches: list[DuoVestiging],
    alle_vestigingen: list[DuoVestiging],
    waarden: dict[str, DuoPersoneelswaarde],
    bron_url: str,
) -> SourceDocument | None:
    if not matches:
        return None
    codes = sorted({item.instellingscode for item in matches})
    gevonden = {code: waarden[code] for code in codes if code in waarden}
    if not gevonden:
        return None
    jaren = {item.jaar for item in gevonden.values()}
    jaar = max(jaren)
    waarden_per_instelling = {
        code: item.aantal_personen for code, item in gevonden.items()
    }
    enkelvoudig = len(codes) == 1
    wp_gevonden = (
        next(iter(gevonden.values())).aantal_personen
        if enkelvoudig else None
    )
    details = ", ".join(
        f"{code}: {waarde.aantal_personen if waarde.aantal_personen is not None else waarde.statistische_markering or 'geen waarde'}"
        for code, waarde in gevonden.items()
    )
    bewijs = (
        f"DUO registreert voor instellingscode {details} onderwijspersoneel "
        f"in aantal personen op 1 oktober {jaar}."
    )
    if not enkelvoudig:
        bewijs += (
            " De waarden zijn niet opgeteld, omdat personen bij meerdere "
            "instellingen dubbel kunnen voorkomen."
        )
    return SourceDocument(
        naam=context.naam,
        company_website_url=context.website_url,
        url=bron_url,
        titel=f"DUO onderwijspersoneel in personen — {context.naam}",
        tekst=bewijs,
        brontype="overheid",
        documenttype="duo_personeel_personen",
        gevraagd_jaar=context.gevraagd_jaar,
        verslagjaar=jaar,
        informatie_peilmoment=f"1 oktober {jaar}",
        wp_gevonden=wp_gevonden,
        eenheid="onderwijspersoneel_personen" if wp_gevonden is not None else None,
        bewijsfragment=bewijs,
        scope_class=_scope(matches[0].sector, matches, alle_vestigingen),
        research_route="duo",
        raw_data={
            "sector": matches[0].sector,
            "instellingscodes": codes,
            "waarden_per_instelling": waarden_per_instelling,
            "brondefinitie": "onderwijspersoneel met betrekkingsomvang groter dan 0 fte",
        },
    )


def _download_sync(url: str) -> bytes:
    if url not in _DOWNLOAD_CACHE:
        response = requests.get(
            url,
            timeout=90,
            headers={"User-Agent": "KYC4etil/1.0 (+openbare-bronnenresearch)"},
            stream=True,
        )
        response.raise_for_status()
        lengte = int(response.headers.get("content-length") or 0)
        if lengte > _MAX_DOWNLOAD_BYTES:
            raise ValueError("DUO-bestand is groter dan de veilige downloadlimiet")
        chunks = []
        totaal = 0
        for chunk in response.iter_content(1024 * 1024):
            totaal += len(chunk)
            if totaal > _MAX_DOWNLOAD_BYTES:
                raise ValueError("DUO-bestand is groter dan de veilige downloadlimiet")
            chunks.append(chunk)
        _DOWNLOAD_CACHE[url] = b"".join(chunks)
    return _DOWNLOAD_CACHE[url]


async def _personeels_url(
    sector: Sector,
) -> str:
    pagina = _PERSONEEL_PAGINAS[sector]
    content = (await asyncio.to_thread(_download_sync, pagina)).decode(
        "utf-8", "ignore",
    )
    hrefs = re.findall(r'href=["\']([^"\']+)["\']', content, re.I)
    marker = _PERSONEEL_MARKER[sector]
    for href in hrefs:
        decoded = html.unescape(href)
        if marker in decoded.lower() and decoded.lower().endswith(".xlsx"):
            return urljoin(pagina, decoded)
    raise ValueError(f"DUO-personeelsbestand voor {sector} niet gevonden")


async def vind_duo_personeelsbron(
    context: QueryContext,
) -> SourceDocument | None:
    """Download de actuele DUO-bronnen en koppel ze aan één organisatie."""
    async def laad_adressen(sector: Sector):
        if sector not in _VESTIGING_CACHE:
            data = await asyncio.to_thread(_download_sync, _ADRES_URLS[sector])
            _VESTIGING_CACHE[sector] = await asyncio.to_thread(
                lees_vestigingen, data, sector,
            )
        return _VESTIGING_CACHE[sector]

    adresgroepen = await asyncio.gather(*[
        laad_adressen(sector) for sector in ("po", "vo", "mbo")
    ])
    alle = [item for groep in adresgroepen for item in groep]
    matches = vind_vestigingen(context, alle)
    if not matches:
        return None
    sector = matches[0].sector
    matches = [item for item in matches if item.sector == sector]
    codes = {item.instellingscode for item in matches}
    bron_url = await _personeels_url(sector)
    bestand = await asyncio.to_thread(_download_sync, bron_url)
    waarden = await asyncio.to_thread(
        lees_personeelswaarden,
        bestand,
        codes,
        context.gevraagd_jaar,
    )
    return bouw_duo_bron(
        context,
        matches,
        _VESTIGING_CACHE[sector],
        waarden,
        bron_url,
    )
