"""Batch-API voor de bronnenwerkbank.

De actieve workflow kent alleen organisaties, research-runs en bronkandidaten.
Oude tabellen worden bij het verwijderen van historische batches nog defensief
opgeruimd, maar spelen geen rol meer in lezen, onderzoeken of exporteren.
"""
import asyncio
import csv
import io
from datetime import datetime, timezone
from pathlib import Path
from zipfile import BadZipFile

from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel, Field
from sqlalchemy import func
from sqlalchemy.orm import Session

from ..auth import get_current_user
from .. import handelingen
from ..database import get_db
from ..models import (
    AgentResult, Batch, BronKandidaat, CallListItem, Candidate, ChatSession,
    Company, Enrichment, JaarverslagMonitoring, PipelineRun, ResearchRun, User,
    VastgoedRecord, WPRecord,
)
from ..research.kostenraming import kosten_per_organisatie
from ..research.service import (
    companies_zonder_afgeronde_research,
    run_research_batch,
)

router = APIRouter(
    prefix="/batches", tags=["batches"], dependencies=[Depends(get_current_user)],
)

CSV_VELDEN = {"naam"}
# Spelling verschilt per export; de VVL-lijsten en de handmatige Excels noemen
# de SBI-tekst allebei anders. "omschrijving" zonder voorvoegsel mag hier op
# sbi_omschrijving uitkomen: het is de enige omschrijvende kolom in dit formaat.
KOLOM_ALIASSEN = {
    "vestnr": "vestigingsnummer",
    "omschrijving": "sbi_omschrijving",
    "sbi omschrijving": "sbi_omschrijving",
    "sbi-omschrijving": "sbi_omschrijving",
    "sbi_code_omschrijving": "sbi_omschrijving",
}
MONITORINGLIJST_NAAM_MARKERS = ("monitoringlijst", "monitorlijst", "watchlist")


def _lijkt_monitoringlijst_batch(batch: Batch) -> bool:
    naam = (batch.naam or "").lower()
    return any(marker in naam for marker in MONITORINGLIJST_NAAM_MARKERS)


def _latest_research_maps(
    db: Session,
    batch_id: str,
) -> tuple[dict[str, ResearchRun], dict[str, BronKandidaat], set[str]]:
    latest_runs: dict[str, ResearchRun] = {}
    for run in (
        db.query(ResearchRun)
        .filter_by(batch_id=batch_id)
        .order_by(ResearchRun.created_at.desc())
    ):
        latest_runs.setdefault(run.company_id, run)

    run_ids = [run.id for run in latest_runs.values()]
    top_candidates: dict[str, BronKandidaat] = {}
    accepted_run_ids: set[str] = set()
    if run_ids:
        candidates = (
            db.query(BronKandidaat)
            .filter(BronKandidaat.research_run_id.in_(run_ids))
            .order_by(BronKandidaat.rang, BronKandidaat.created_at)
            .all()
        )
        for candidate in candidates:
            top_candidates.setdefault(candidate.research_run_id, candidate)
            if candidate.status == "geaccepteerd":
                accepted_run_ids.add(candidate.research_run_id)
                top_candidates[candidate.research_run_id] = candidate
    return latest_runs, top_candidates, accepted_run_ids


def _run_research_batch_background(batch_id: str) -> None:
    asyncio.run(run_research_batch(batch_id))


def _lees_upload(content: bytes, bestandsnaam: str) -> tuple[set[str], list[dict]]:
    if Path(bestandsnaam).suffix.lower() == ".xlsx":
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            worksheet = workbook.active
            waarden = worksheet.iter_rows(values_only=True)
            kop = next(waarden, ())
            rows = [dict(zip(kop, row)) for row in waarden]
            workbook.close()
        except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
            raise HTTPException(422, "Ongeldig Excel-bestand") from exc
    else:
        try:
            tekst = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise HTTPException(422, "Bestand moet CSV of Excel (.xlsx) zijn") from exc
        reader = csv.DictReader(io.StringIO(tekst))
        kop = reader.fieldnames or ()
        rows = list(reader)

    kolommen = {
        KOLOM_ALIASSEN.get(str(k).strip().lower(), str(k).strip().lower())
        for k in kop if k is not None
    }
    genormaliseerd = []
    for row in rows:
        schoon = {
            KOLOM_ALIASSEN.get(str(k).strip().lower(), str(k).strip().lower()):
                (str(v).strip() if v is not None and str(v).strip() else None)
            for k, v in row.items() if k is not None
        }
        if schoon.get("naam"):
            genormaliseerd.append(schoon)
    return kolommen, genormaliseerd


# De velden die uit een bestand of formulier op een Company terechtkomen. Eén
# lijst, zodat de upload, de samenvoeging en het handmatig toevoegen niet elk
# hun eigen selectie krijgen — dat is precies hoe `sbi_omschrijving` jarenlang
# stil kon verdwijnen.
COMPANY_VELDEN = (
    "vestigingsnummer", "naam", "gemeente", "adres", "sbi_code",
    "sbi_omschrijving", "cb_er", "kvk_nummer", "website_url", "telefoonnummer",
)


def _identiteit(velden: dict) -> tuple:
    """Waaraan herken je dat twee rijen dezelfde vestiging zijn?

    Het vestigingsnummer is de identificatie richting VVL en dus het hardst.
    Daarna het KvK-nummer; pas als beide ontbreken naam plus gemeente, want een
    naam alleen komt in meerdere gemeenten voor ("Jumbo Supermarkten").
    """
    if velden.get("vestigingsnummer"):
        return ("vestigingsnummer", str(velden["vestigingsnummer"]).strip().lower())
    if velden.get("kvk_nummer"):
        return ("kvk_nummer", str(velden["kvk_nummer"]).strip().lower())
    return (
        "naam",
        str(velden.get("naam") or "").strip().lower(),
        str(velden.get("gemeente") or "").strip().lower(),
    )


def _vul_lege_velden_aan(company: Company, velden: dict) -> bool:
    """Vul alleen wat nog leeg is. Geeft terug of er iets is veranderd.

    Bewust niet overschrijven: wat er staat kan met de hand zijn gecorrigeerd of
    tijdens een run zijn gevonden, en een nieuwe aanlevering is niet vanzelf
    beter. Wat ontbreekt aanvullen is winst zonder risico.
    """
    gewijzigd = False
    for veld in COMPANY_VELDEN:
        if veld == "naam":
            continue
        nieuw = velden.get(veld)
        if nieuw and not getattr(company, veld, None):
            setattr(company, veld, nieuw)
            gewijzigd = True
    return gewijzigd


def _voeg_rijen_toe(db: Session, batch: Batch, rows: list[dict]) -> dict:
    """Voeg rijen toe aan een bestaande lijst zonder dubbelen te maken."""
    bestaand = {
        _identiteit({
            veld: getattr(company, veld) for veld in COMPANY_VELDEN
        }): company
        for company in db.query(Company).filter_by(batch_id=batch.id)
    }
    toegevoegd = bijgewerkt = 0
    for row in rows:
        company = bestaand.get(_identiteit(row))
        if company is None:
            company = Company(
                batch_id=batch.id,
                **{veld: row.get(veld) for veld in COMPANY_VELDEN},
            )
            db.add(company)
            bestaand[_identiteit(row)] = company
            toegevoegd += 1
        elif _vul_lege_velden_aan(company, row):
            bijgewerkt += 1
    batch.totaal = (batch.totaal or 0) + toegevoegd
    return {
        "toegevoegd": toegevoegd,
        "bijgewerkt": bijgewerkt,
        "ongewijzigd": len(rows) - toegevoegd - bijgewerkt,
    }


@router.post("/upload")
async def upload_batch(
    file: UploadFile,
    naam: str | None = None,
    jaar: int | None = None,
    monitoringlijst: bool = False,
    map_id: str | None = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Maak een onderzoekslijst uit een CSV- of Excel-bestand."""
    kolommen, rows = _lees_upload(await file.read(), file.filename or "")
    if not rows or not CSV_VELDEN.issubset(kolommen):
        raise HTTPException(422, "Bestand mist verplichte kolom 'naam'")

    # Een monitoringlijst is er één, en die loopt door. Een nieuw bestand vulde
    # hem tot nu toe niet aan maar verving hem: de oude lijst raakte stil zijn
    # vlag kwijt en alle controlegeschiedenis verdween uit beeld. Aanvullen is
    # wat je wilt — nieuwe organisaties erbij, bestaande met rust gelaten.
    lopende_watchlist = (
        db.query(Batch).filter_by(is_monitoringlijst=True).first()
        if monitoringlijst else None
    )
    if lopende_watchlist is not None:
        telling = _voeg_rijen_toe(db, lopende_watchlist, rows)
        handelingen.leg_vast(
            db, handelingen.LIJST_GEUPLOAD,
            f"Monitoringlijst aangevuld met {telling['toegevoegd']} organisaties "
            f"uit {file.filename}",
            door=current_user, onderwerp_id=lopende_watchlist.id,
        )
        db.commit()
        return {
            "batch_id": lopende_watchlist.id,
            "aantal_companies": telling["toegevoegd"],
            "samengevoegd": True,
            **telling,
        }

    batch = Batch(
        naam=naam or file.filename,
        jaar=jaar or datetime.now(timezone.utc).year,
        totaal=0,
        is_monitoringlijst=monitoringlijst,
        geupload_door=current_user.id,
        # Een monitoringlijst hoort niet in een map: die heeft een eigen
        # module en zou anders tussen de onderzoekslijsten opduiken.
        map_id=None if monitoringlijst else map_id,
    )
    db.add(batch)
    db.flush()
    telling = _voeg_rijen_toe(db, batch, rows)
    handelingen.leg_vast(
        db, handelingen.LIJST_GEUPLOAD,
        f"Lijst '{batch.naam}' geüpload met {telling['toegevoegd']} organisaties",
        door=current_user, onderwerp_id=batch.id,
    )
    db.commit()
    return {
        "batch_id": batch.id,
        "aantal_companies": telling["toegevoegd"],
        "samengevoegd": False,
        **telling,
    }


class HandmatigBedrijf(BaseModel):
    """Eén organisatie die een reviewer zelf toevoegt."""

    naam: str = Field(min_length=2, max_length=255)
    vestigingsnummer: str | None = Field(default=None, max_length=20)
    gemeente: str | None = Field(default=None, max_length=100)
    adres: str | None = Field(default=None, max_length=500)
    sbi_code: str | None = Field(default=None, max_length=10)
    sbi_omschrijving: str | None = Field(default=None, max_length=500)
    kvk_nummer: str | None = Field(default=None, max_length=20)
    website_url: str | None = Field(default=None, max_length=500)


@router.post("/{batch_id}/companies")
def voeg_bedrijf_toe(
    batch_id: str,
    payload: HandmatigBedrijf,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Voeg één organisatie toe aan een bestaande lijst.

    Dezelfde samenvoeglogica als de upload: bestaat de organisatie al, dan wordt
    ze niet gedupliceerd maar hooguit aangevuld. Zo levert twee keer toevoegen
    niet twee rijen op die daarna allebei onderzocht worden.
    """
    batch = db.get(Batch, batch_id)
    if batch is None:
        raise HTTPException(404, "lijst niet gevonden")

    velden = payload.model_dump()
    telling = _voeg_rijen_toe(db, batch, [velden])
    if telling["toegevoegd"]:
        handelingen.leg_vast(
            db, handelingen.ORGANISATIE_TOEGEVOEGD,
            f"{payload.naam} met de hand toegevoegd aan '{batch.naam}'",
            door=current_user, onderwerp_id=batch.id,
        )
    db.commit()
    company = (
        db.query(Company)
        .filter_by(batch_id=batch.id, naam=payload.naam)
        .order_by(Company.created_at.desc())
        .first()
    )
    return {
        "company_id": company.id if company else None,
        "bestond_al": telling["toegevoegd"] == 0,
        **telling,
    }


@router.post("/{batch_id}/run")
async def start_batch(
    batch_id: str,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
):
    """Start autonome bronnenresearch voor de volledige lijst."""
    batch = db.get(Batch, batch_id)
    if batch is None:
        raise HTTPException(404, f"batch {batch_id} bestaat niet")
    if batch.status == "running":
        return {"batch_id": batch.id, "status": batch.status,
                "verwerkt": batch.verwerkt, "totaal": batch.totaal}

    afgerond = (
        db.query(func.count(func.distinct(ResearchRun.company_id)))
        .filter_by(batch_id=batch.id, status="completed")
        .scalar() or 0
    )
    batch.status = "running"
    batch.verwerkt = afgerond
    batch.completed_at = None
    db.commit()
    background_tasks.add_task(_run_research_batch_background, batch.id)
    return {"batch_id": batch.id, "status": batch.status,
            "verwerkt": batch.verwerkt, "totaal": batch.totaal,
            "workflow": "autonome_bronnenresearch"}


@router.post("/{batch_id}/cancel")
def cancel_batch(batch_id: str, db: Session = Depends(get_db)):
    batch = db.get(Batch, batch_id)
    if batch is None:
        raise HTTPException(404, "batch niet gevonden")
    if batch.status != "running":
        raise HTTPException(409, f"batch heeft status '{batch.status}', niet 'running'")
    batch.status = "cancelled"
    db.commit()
    return {"batch_id": batch.id, "status": batch.status}


@router.post("/{batch_id}/reset-vastgelopen")
def reset_vastgelopen(batch_id: str, db: Session = Depends(get_db)):
    batch = db.get(Batch, batch_id)
    if batch is None:
        raise HTTPException(404, "batch niet gevonden")
    if batch.status != "running":
        raise HTTPException(409, "alleen een lopende batch kan worden gereset")
    batch.status = "pending"
    db.commit()
    return {"batch_id": batch.id, "status": batch.status}


@router.delete("/{batch_id}")
def delete_batch(
    batch_id: str,
    definitief: bool = False,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Gooi een lijst weg — standaard naar de prullenbak.

    Verwijderen wiste eerder alles ineens: elke organisatie, elke bronkandidaat,
    elke beoordeling die erin zat, zonder weg terug. In een werkbank waar het
    werk juist in die beoordelingen zit is dat te scherp. Nu verdwijnt de lijst
    uit beeld en is ze terug te halen; `definitief=true` wist haar echt.
    """
    batch = db.get(Batch, batch_id)
    if batch is None:
        raise HTTPException(404, "batch niet gevonden")
    if batch.status == "running":
        raise HTTPException(409, "batch draait nog; annuleer eerst")

    if not definitief:
        batch.verwijderd_op = datetime.now(timezone.utc).replace(tzinfo=None)
        handelingen.leg_vast(
            db, handelingen.LIJST_VERWIJDERD,
            f"Lijst '{batch.naam}' naar de prullenbak",
            door=current_user, onderwerp_id=batch.id,
        )
        db.commit()
        return {"deleted": batch_id, "definitief": False}

    handelingen.leg_vast(
        db, handelingen.LIJST_VERWIJDERD,
        f"Lijst '{batch.naam}' definitief verwijderd, met alles wat erin zat",
        door=current_user, onderwerp_id=batch.id,
    )
    company_ids = [row[0] for row in db.query(Company.id).filter_by(batch_id=batch_id)]
    if company_ids:
        for model in (CallListItem, WPRecord, ChatSession, Enrichment,
                      JaarverslagMonitoring, VastgoedRecord):
            db.query(model).filter(model.company_id.in_(company_ids)).delete(
                synchronize_session=False,
            )
        db.query(BronKandidaat).filter(
            BronKandidaat.company_id.in_(company_ids),
        ).delete(synchronize_session=False)
        db.query(ResearchRun).filter(
            ResearchRun.company_id.in_(company_ids),
        ).delete(synchronize_session=False)
    for model in (Candidate, AgentResult, PipelineRun):
        db.query(model).filter(model.batch_id == batch_id).delete(
            synchronize_session=False,
        )
    db.query(Company).filter_by(batch_id=batch_id).delete(synchronize_session=False)
    db.delete(batch)
    db.commit()
    return {"deleted": batch_id, "definitief": True}


@router.post("/{batch_id}/herstel")
def herstel_batch(
    batch_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Haal een weggegooide lijst terug uit de prullenbak."""
    batch = db.get(Batch, batch_id)
    if batch is None:
        raise HTTPException(404, "batch niet gevonden")
    batch.verwijderd_op = None
    handelingen.leg_vast(
        db, handelingen.LIJST_HERSTELD,
        f"Lijst '{batch.naam}' teruggehaald uit de prullenbak",
        door=current_user, onderwerp_id=batch.id,
    )
    db.commit()
    return {"hersteld": batch_id}


@router.get("")
def list_batches(
    map_id: str | None = None,
    losse_lijsten: bool = False,
    prullenbak: bool = False,
    db: Session = Depends(get_db),
):
    """Lijsten, optioneel beperkt tot één map.

    Zonder filter blijft dit alle lijsten teruggeven, zodat bestaande
    aanroepers niet veranderen. `losse_lijsten=true` geeft juist alleen wat
    buiten elke map staat.
    """
    query = db.query(Batch).filter(Batch.is_monitoringlijst.isnot(True))
    # Weggegooide lijsten staan standaard niet tussen de rest; `prullenbak`
    # laat juist alleen die zien.
    query = (
        query.filter(Batch.verwijderd_op.isnot(None)) if prullenbak
        else query.filter(Batch.verwijderd_op.is_(None))
    )
    if map_id:
        query = query.filter(Batch.map_id == map_id)
    elif losse_lijsten:
        query = query.filter(Batch.map_id.is_(None))
    batches = [
        batch for batch in query.order_by(Batch.created_at.desc()).all()
        if not _lijkt_monitoringlijst_batch(batch)
    ]
    # Hoeveel er nog te onderzoeken zijn, met dezelfde regel als de run zelf.
    # Zonder dit rekende de kostenindicatie op het scherm met de hele lijst,
    # ook als er nog maar een handvol organisaties over was.
    nog_te_doen = {
        batch.id: len(companies_zonder_afgeronde_research(db, batch.id))
        for batch in batches
    }
    uploader_ids = {batch.geupload_door for batch in batches if batch.geupload_door}
    naam_per_id = {
        user.id: user.naam
        for user in db.query(User).filter(User.id.in_(uploader_ids)).all()
    } if uploader_ids else {}
    return [{
        "id": batch.id,
        "naam": batch.naam,
        "jaar": batch.jaar,
        "status": batch.status,
        "totaal": batch.totaal,
        "verwerkt": batch.verwerkt,
        "created_at": batch.created_at.isoformat() + "Z" if batch.created_at else None,
        "completed_at": (
            batch.completed_at.isoformat() + "Z" if batch.completed_at else None
        ),
        "geupload_door_naam": naam_per_id.get(batch.geupload_door),
        "map_id": batch.map_id,
        "nog_te_onderzoeken": nog_te_doen[batch.id],
    } for batch in batches]


@router.get("/kostenindicatie")
def get_kostenindicatie(db: Session = Depends(get_db)):
    """Wat kost het onderzoeken van één organisatie, volgens de vorige runs."""
    return kosten_per_organisatie(db)


@router.get("/{batch_id}")
def get_batch(batch_id: str, db: Session = Depends(get_db)):
    batch = db.get(Batch, batch_id)
    if batch is None:
        raise HTTPException(404, "batch niet gevonden")
    latest_runs, top_candidates, accepted_run_ids = _latest_research_maps(db, batch_id)
    research = {
        "gestart": len(latest_runs),
        "afgerond": sum(run.status == "completed" for run in latest_runs.values()),
        "review_nodig": sum(
            run.status == "completed"
            and run.id not in accepted_run_ids
            and run.id in top_candidates
            for run in latest_runs.values()
        ),
        "geaccepteerd": len(accepted_run_ids),
        "niet_gevonden": sum(
            run.status == "completed" and run.resultaat_status == "niet_gevonden"
            for run in latest_runs.values()
        ),
        "fouten": sum(run.status == "error" for run in latest_runs.values()),
    }
    return {
        "id": batch.id, "naam": batch.naam, "jaar": batch.jaar,
        "status": batch.status, "totaal": batch.totaal,
        "verwerkt": batch.verwerkt, "workflow": "autonome_bronnenresearch",
        "research": research,
        "created_at": batch.created_at.isoformat() + "Z" if batch.created_at else None,
        "completed_at": (
            batch.completed_at.isoformat() + "Z" if batch.completed_at else None
        ),
    }


@router.get("/{batch_id}/companies")
def list_companies(batch_id: str, db: Session = Depends(get_db)):
    if db.get(Batch, batch_id) is None:
        raise HTTPException(404, "batch niet gevonden")
    latest_runs, top_candidates, accepted_run_ids = _latest_research_maps(db, batch_id)
    out = []
    for company in db.query(Company).filter_by(batch_id=batch_id).order_by(Company.naam):
        run = latest_runs.get(company.id)
        top = top_candidates.get(run.id) if run else None
        wp_bruikbaar = bool(
            top and top.eenheid == "werkzame_personen"
            and top.scope_class in {"vestiging", "limburg"}
        )
        out.append({
            "company_id": company.id,
            "naam": company.naam,
            "gemeente": company.gemeente,
            "vestigingsnummer": company.vestigingsnummer,
            "cb_er": company.cb_er,
            "kvk_nummer": company.kvk_nummer,
            "sbi_omschrijving": company.sbi_omschrijving,
            "research_run_id": run.id if run else None,
            "research_status": run.status if run else "niet_gestart",
            "research_resultaat_status": run.resultaat_status if run else None,
            "research_top_id": top.id if top else None,
            "research_top_url": top.url if top else None,
            "research_top_titel": top.titel if top else None,
            "research_top_type": top.brontype if top else None,
            "research_top_wp": top.wp_gevonden if wp_bruikbaar else None,
            "research_top_wp_raw": top.wp_gevonden if top else None,
            "research_top_wp_bruikbaar": wp_bruikbaar,
            "research_top_scope": top.scope_class if top else None,
            "research_top_score": top.ranking_score if top else None,
            "research_review_status": (
                "geaccepteerd" if run and run.id in accepted_run_ids
                else top.status if top else None
            ),
        })
    return out


@router.get("/{batch_id}/export.xlsx")
def export_bronnen(batch_id: str, db: Session = Depends(get_db)):
    """Exporteer per organisatie alleen de door een mens gekozen bron."""
    batch = db.get(Batch, batch_id)
    if batch is None:
        raise HTTPException(404, "batch niet gevonden")
    latest_runs, gekozen_per_run, _ = _latest_research_maps(db, batch_id)
    reviewer_ids = {
        bron.reviewed_by for bron in gekozen_per_run.values() if bron.reviewed_by
    }
    reviewer_namen = {
        user.id: user.naam
        for user in db.query(User).filter(User.id.in_(reviewer_ids)).all()
    } if reviewer_ids else {}
    headers = [
        "Vestigingsnummer", "Organisatienaam", "Gemeente", "KvK-nummer",
        "Website", "Batchjaar", "Researchstatus", "Reviewstatus",
        "Bron-URL", "Brontype",
        "Documenttype", "Verslagjaar", "Informatiepeilmoment", "Gevonden waarde",
        "Eenheid", "Scope", "Identity class", "Bewijsfragment", "PDF-pagina",
        "Ranking-score", "Reviewredencode", "Reviewtoelichting", "Beoordeeld door",
        "Beoordeeld op",
    ]
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Gekozen bronnen"
    worksheet.append(headers)
    for company in db.query(Company).filter_by(batch_id=batch_id).order_by(Company.naam):
        run = latest_runs.get(company.id)
        bron = gekozen_per_run.get(run.id) if run else None
        if bron is not None and bron.status != "geaccepteerd":
            bron = None
        worksheet.append([
            company.vestigingsnummer, company.naam, company.gemeente,
            company.kvk_nummer,
            company.effectieve_website_url,
            batch.jaar,
            run.status if run else "niet_gestart",
            bron.status if bron else None,
            bron.url if bron else None,
            bron.brontype if bron else None,
            bron.documenttype if bron else None,
            bron.verslagjaar if bron else None,
            bron.informatie_peilmoment if bron else None,
            bron.wp_gevonden if bron else None,
            bron.eenheid if bron else None,
            bron.scope_class if bron else None,
            bron.identity_class if bron else None,
            bron.bewijsfragment if bron else None,
            bron.bron_pagina if bron else None,
            bron.ranking_score if bron else None,
            bron.review_reason_code if bron else None,
            bron.review_reason if bron else None,
            reviewer_namen.get(bron.reviewed_by) if bron else None,
            bron.reviewed_at.isoformat() + "Z" if bron and bron.reviewed_at else None,
        ])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    buffer = io.BytesIO()
    workbook.save(buffer)
    filename = f"bronnen_{batch.naam or batch.id}.xlsx".replace('"', "")
    return StreamingResponse(
        iter([buffer.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
